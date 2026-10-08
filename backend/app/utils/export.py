"""
backend/app/utils/export.py

3-Mode Event Export Generator (CSV & Excel XLSX) with DDE Formula Injection Sanitization.
Milestone 4: Feature 24.
"""

from __future__ import annotations

import csv
import io
import time
from typing import Any, Dict, List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


EXPORT_HEADERS = ["Date Time", "Camera Name", "Event Type", "Description", "Snapshot Link"]


def sanitize_spreadsheet_cell(val: Any) -> str:
    """
    Sanitize cell value to defend against Dynamic Data Exchange (DDE) formula injection.
    Prepends single quote (') if string starts with '=', '+', '-', or '@'.
    """
    if val is None:
        return ""
    s = str(val).strip()
    if s and s[0] in ("=", "+", "-", "@"):
        return "'" + s
    return str(val)


def format_export_timestamp(ts: Any) -> str:
    """Format numeric epoch or ISO timestamp for export output."""
    if ts is None:
        return ""
    if isinstance(ts, (int, float)):
        return time.strftime("%Y-%m-%d %H:%M:%S", time.gmtime(float(ts)))
    s = str(ts).strip()
    try:
        val = float(s)
        return time.strftime("%Y-%m-%d %H:%M:%S", time.gmtime(val))
    except ValueError:
        return s


def export_events_csv(records: List[Dict[str, Any]], mode: str = "all") -> str:
    """
    Generate CSV string for 3-mode export:
    - mode="template": Strictly header-only row (no data rows).
    - mode="filtered": Header row + sanitized filtered data rows.
    - mode="all" (or unknown): Header row + all sanitized data rows.
    """
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL, lineterminator="\n")

    # 1. Write header row
    writer.writerow(EXPORT_HEADERS)

    # 2. Template mode strictly outputs header row only
    if mode == "template":
        return output.getvalue()

    # 3. Write data rows
    for r in records:
        dt = sanitize_spreadsheet_cell(format_export_timestamp(r.get("timestamp")))
        cam_name = sanitize_spreadsheet_cell(r.get("camera_name", ""))
        evt_type = sanitize_spreadsheet_cell(r.get("event_type", ""))
        desc = sanitize_spreadsheet_cell(r.get("description", ""))
        snap = sanitize_spreadsheet_cell(r.get("snapshot_url") or r.get("snapshot_path", ""))
        writer.writerow([dt, cam_name, evt_type, desc, snap])

    return output.getvalue()


def export_events_xlsx(records: List[Dict[str, Any]], mode: str = "all") -> bytes:
    """
    Generate genuine Excel .xlsx binary buffer with styled header and auto-width columns.
    - mode="template": Strictly header row only (no data rows).
    - mode="filtered": Header row + sanitized filtered data rows.
    - mode="all" (or unknown): Header row + all sanitized data rows.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Event Logs"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    regular_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )

    # 1. Write header row
    ws.append(EXPORT_HEADERS)
    for col_idx in range(1, len(EXPORT_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")

    # 2. Write data rows if not template mode
    if mode != "template":
        for r in records:
            dt = sanitize_spreadsheet_cell(format_export_timestamp(r.get("timestamp")))
            cam_name = sanitize_spreadsheet_cell(r.get("camera_name", ""))
            evt_type = sanitize_spreadsheet_cell(r.get("event_type", ""))
            desc = sanitize_spreadsheet_cell(r.get("description", ""))
            snap = sanitize_spreadsheet_cell(r.get("snapshot_url") or r.get("snapshot_path", ""))
            row_data = [dt, cam_name, evt_type, desc, snap]
            ws.append(row_data)

            curr_row = ws.max_row
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
